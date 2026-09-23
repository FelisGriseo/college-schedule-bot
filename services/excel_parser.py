import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

from services.text_utils import match_group_name, normalize_group_name

REQUIRED_COLUMNS = {
    "course", "group", "day", "pair_number", "pair_time", "subject", "teacher", "room"
}


def _parse_pair_number(value: object) -> int:
    """Accept Excel values such as 1, 1.0 and a lesson range like 1-2."""
    text = str(value).strip()
    match = re.match(r"^(\d+)(?:\s*[-–]\s*\d+)?(?:\.0)?$", text)
    if not match:
        raise ValueError(f"Unsupported pair_number value in Excel: {value!r}")
    return int(match.group(1))


@dataclass(frozen=True)
class ScheduleItem:
    course: str
    group: str
    day: str
    pair_number: int
    pair_time: str
    subject: str
    teacher: str
    room: str
    week_type: str = "all"
    pair_label: str = ""
    subgroup: str = "Вся група"
    lesson_type: str = ""


class ExcelSchedule:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.items: list[ScheduleItem] = []
        self._group_names: set[str] = set()
        self._load()

    def _load(self) -> None:
        if not self.path.exists():
            raise FileNotFoundError(
                f"Schedule file not found: {self.path}. Put the Excel schedule in the project root."
            )
        workbook = pd.ExcelFile(self.path)
        if "Parsed_Schedule" in workbook.sheet_names:
            self._load_flat(workbook)
        else:
            self._load_college_layout()

    def _load_flat(self, workbook: pd.ExcelFile) -> None:
        frame = pd.read_excel(workbook, sheet_name="Parsed_Schedule")
        frame.columns = [str(column).strip().lower() for column in frame.columns]
        missing = REQUIRED_COLUMNS - set(frame.columns)
        if missing:
            raise ValueError(f"Parsed_Schedule is missing columns: {', '.join(sorted(missing))}")

        for row in frame.fillna("").to_dict(orient="records"):
            if not str(row["group"]).strip() or not str(row["subject"]).strip():
                continue
            self.items.append(ScheduleItem(
                course=str(row["course"]).strip(),
                group=normalize_group_name(str(row["group"])),
                day=str(row["day"]).strip(),
                pair_number=_parse_pair_number(row["pair_number"]),
                pair_time=str(row["pair_time"]).strip(),
                subject=str(row["subject"]).strip(),
                teacher=str(row["teacher"]).strip(),
                room=str(row["room"]).strip(),
                pair_label=str(row["pair_number"]).strip(),
            ))
            self._group_names.add(normalize_group_name(str(row["group"])))

    def _load_college_layout(self) -> None:
        workbook = load_workbook(self.path, data_only=True, read_only=False)
        day_names = {
            "понеділок": "monday",
            "вівторок": "tuesday",
            "середа": "wednesday",
            "четвер": "thursday",
            "п\'ятниця": "friday",
            "п’ятниця": "friday",
            "субота": "saturday",
            "неділя": "sunday",
        }
        group_pattern = re.compile(r"^[A-Za-zА-Яа-яІіЇїЄєҐґ]{1,6}-\d{3}$")
        pair_pattern = re.compile(r"^\d+\s*[-–]\s*\d+$")

        for worksheet in workbook.worksheets:
            sheet_title = worksheet.title.strip().lower()
            if not re.match(r"^\d+\s*курс$", sheet_title) and sheet_title != "бакалавр":
                continue
            groups = []
            for row in worksheet.iter_rows(min_row=1, max_row=min(15, worksheet.max_row)):
                for cell in row:
                    value = str(cell.value).strip() if cell.value is not None else ""
                    if group_pattern.fullmatch(value):
                        groups.append((value, cell.column))
            if not groups:
                continue
            self._group_names.update(normalize_group_name(group) for group, _ in groups)

            day_rows = []
            for row_number in range(1, worksheet.max_row + 1):
                value = worksheet.cell(row_number, 1).value
                normalized = str(value).strip().lower() if value is not None else ""
                if normalized in day_names:
                    day_rows.append((row_number, normalized))

            for day_index, (day_row, day_name) in enumerate(day_rows):
                end_row = day_rows[day_index + 1][0] if day_index + 1 < len(day_rows) else worksheet.max_row + 1
                pair_rows = [
                    row_number
                    for row_number in range(day_row, end_row)
                    if pair_pattern.fullmatch(str(worksheet.cell(row_number, 2).value).strip())
                ]
                for pair_index, pair_row in enumerate(pair_rows):
                    block_end = pair_rows[pair_index + 1] if pair_index + 1 < len(pair_rows) else end_row
                    pair_label = str(worksheet.cell(pair_row, 2).value).strip()
                    pair_number = _parse_pair_number(pair_label)
                    pair_time = _find_pair_time(worksheet, pair_row, block_end)
                    for group_name, column in groups:
                        slots = _read_group_pair(
                            worksheet, column, pair_row, block_end
                        )
                        for subject, teacher, room, week_type, subgroup in slots:
                            self.items.append(ScheduleItem(
                                course=worksheet.title.strip(),
                                group=normalize_group_name(group_name),
                                day=day_names[day_name],
                                pair_number=pair_number,
                                pair_time=pair_time,
                                subject=subject,
                                teacher=teacher,
                                room=room,
                                week_type=week_type,
                                pair_label=pair_label,
                                subgroup=subgroup,
                                lesson_type=_lesson_type(subject),
                            ))

    def groups_by_course(self) -> dict[str, list[str]]:
        result: dict[str, set[str]] = {}
        for item in self.items:
            normalized_group = normalize_group_name(item.group)
            result.setdefault(item.course, set()).add(normalized_group)
        return {course: sorted(groups) for course, groups in sorted(result.items())}

    @property
    def group_names(self) -> tuple[str, ...]:
        return tuple(sorted(self._group_names))

    def resolve_group_name(self, raw: str) -> tuple[str | None, str, list[str]]:
        return match_group_name(raw, list(self._group_names))

    def for_group_and_date(self, group: str, target: date) -> list[ScheduleItem]:
        wanted_group = normalize_group_name(group)
        day_names = {
            0: ("понеділок", "понедельник", "monday"),
            1: ("вівторок", "вторник", "tuesday"),
            2: ("середа", "среда", "wednesday"),
            3: ("четвер", "четверг", "thursday"),
            4: ("п’ятниця", "п'ятниця", "пятница", "friday"),
            5: ("субота", "суббота", "saturday"),
            6: ("неділя", "воскресенье", "sunday"),
        }
        aliases = day_names[target.weekday()]
        week_type = _week_type_for_date(target)
        return sorted(
            [
                item for item in self.items
                if normalize_group_name(item.group) == wanted_group
                and item.day.strip().lower() in aliases
                and item.week_type in ("all", week_type)
            ],
            key=lambda item: (item.pair_number, item.week_type),
        )


def _find_pair_time(worksheet, start_row: int, end_row: int) -> str:
    values = []
    for row in range(start_row, end_row):
        value = worksheet.cell(row, 3).value
        if isinstance(value, (datetime, time)):
            formatted = value.strftime("%H:%M")
            if formatted not in values:
                values.append(formatted)
    if len(values) >= 2:
        return f"{values[0]} - {values[1]}"
    return values[0] if values else ""


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_group_pair(
    worksheet, column: int, start_row: int, end_row: int
) -> list[tuple[str, str, str, str, str]]:
    """Read subject/teacher/room slots from one group block and one pair."""
    subject_rows = [
        row
        for row in range(start_row, end_row)
        if _is_subject_cell(worksheet.cell(row, column))
    ]
    slots: list[tuple[str, str, str, str, str]] = []
    subject_groups: list[tuple[int, str, list[int]]] = []
    for index, subject_row in enumerate(subject_rows):
        next_subject = subject_rows[index + 1] if index + 1 < len(subject_rows) else end_row
        subject = str(worksheet.cell(subject_row, column).value).strip()
        teacher_rows = [
            row
            for row in range(subject_row + 1, next_subject)
            if worksheet.cell(row, column).value not in (None, "")
            and not _is_subject_cell(worksheet.cell(row, column))
        ]
        subject_groups.append((index, subject, teacher_rows))

    for index, subject, teacher_rows in subject_groups:
        week_type = (
            "all"
            if len(subject_groups) == 1
            else "numerator"
            if index == 0
            else "denominator"
        )
        if not teacher_rows:
            slots.append((subject, "", "", week_type, "Вся група"))
            continue
        subgroup_count = len(teacher_rows)
        for subgroup_index, teacher_row in enumerate(teacher_rows):
            teacher = str(worksheet.cell(teacher_row, column).value).strip()
            room_value = worksheet.cell(teacher_row, column + 2).value
            room = _format_cell_value(room_value)
            subgroup = (
                f"{subgroup_index + 1} підгрупа"
                if subgroup_count > 1
                else "Вся група"
            )
            slots.append((subject, teacher, room, week_type, subgroup))
    return slots


def _is_subject_cell(cell) -> bool:
    value = cell.value
    return isinstance(value, str) and bool(value.strip()) and bool(cell.font.bold)


def _lesson_type(subject: str) -> str:
    lowered = subject.lower()
    if re.search(r"\bлаб(?:ораторна|ораторна|\.)?\b", lowered):
        return "Лабораторна"
    if re.search(r"\bлек(?:ція|ція|\.)?\b", lowered):
        return "Лекція"
    if re.search(r"\b(практ|практична|пр\.)\b", lowered):
        return "Практична"
    if re.search(r"\bсемінар\b", lowered):
        return "Семінар"
    return ""


def _week_type_for_date(target: date) -> str:
    today = datetime.now(ZoneInfo("Europe/Kyiv")).date()
    current_monday = today - timedelta(days=today.weekday())
    target_monday = target - timedelta(days=target.weekday())
    weeks_from_current = (target_monday - current_monday).days // 7
    return "denominator" if weeks_from_current % 2 == 0 else "numerator"
